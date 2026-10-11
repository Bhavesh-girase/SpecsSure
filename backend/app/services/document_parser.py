import io
import logging
from typing import Tuple, Optional

logger = logging.getLogger(__name__)

def extract_text_from_file(file_bytes: bytes, filename: str) -> Tuple[str, Optional[str]]:
    """
    Extracts plain text from uploaded file bytes based on file extension.
    Supports PDF (.pdf), DOCX (.docx), and plain text (.txt, .md, .csv).
    Returns (extracted_text, error_message).
    """
    lower_name = filename.lower()
    try:
        if lower_name.endswith(".pdf"):
            return extract_from_pdf(file_bytes), None
        elif lower_name.endswith(".docx"):
            return extract_from_docx(file_bytes), None
        elif lower_name.endswith(".txt") or lower_name.endswith(".md") or lower_name.endswith(".csv"):
            return file_bytes.decode("utf-8", errors="replace"), None
        else:
            # Attempt general UTF-8 decode
            return file_bytes.decode("utf-8", errors="replace"), None
    except Exception as e:
        logger.error(f"Error extracting text from {filename}: {e}", exc_info=True)
        return "", f"Failed to parse document: {str(e)}"

def extract_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF using PyMuPDF (fitz) or fallback."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                text_parts.append(f"--- Page {page_num + 1} ---\n{text}")
        doc.close()
        return "\n\n".join(text_parts)
    except Exception as e:
        logger.warning(f"PyMuPDF failed, trying pypdf/fallback: {e}")
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = [page.extract_text() for page in reader.pages if page.extract_text()]
            return "\n\n".join(text_parts)
        except Exception as e2:
            raise RuntimeError(f"Could not extract text from PDF: {e2}")

def extract_from_docx(file_bytes: bytes) -> str:
    """Extracts text and tables from DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        
        # Also extract table contents
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(f"| {row_text} |")
                    
        return "\n".join(full_text)
    except Exception as e:
        raise RuntimeError(f"Could not extract text from DOCX: {e}")
