"""
Tender Synthesizer — "Intent → verified tender" in one pass.

    intent (line items + slots)
        │  Procurement Knowledge Graph, as of the publication date
        ▼
    clause plan  ── every sentence carries its basis (standard / regulation / buyer / practice / policy)
        │           and the KG path that produced it (item → family → edition in force → regulation …)
        ▼
    render ──► Tender Compiler (IR → requirements → Z3 → market → 38 lint rules)
        ▲                                   │
        └──── self-repair loop ◀────────────┘  safe autofixes (test methods, NABL, acceptance, certification)
                                               + Competition Guard (relax 'practice' defaults, never buyer choices,
                                                 until ≥ 3 manufacturers can comply)

The generator is checked by the same compiler that checks human-written tenders: it cannot hand over a draft
with an open critical finding without saying so.
"""
from __future__ import annotations

import hashlib
import io
import re
from datetime import date, timedelta
from typing import Dict, List, Optional, Tuple

from app.engine import ontology
from app.engine.compiler import compile_ir
from app.engine.docparse import parse_text
from app.engine.kb import KnowledgeBase
from app.engine.kg import ProcurementKG, get_kg, item_by_id, policies
from app.engine.market import MIN_COMPETITION, catalog
from app.engine.std_refs import family_of
from app.engine.tender_ir import build_ir

GENERATOR_VERSION = "tender-synthesizer/1.0"
SAFE_FIX_RULES = {"SS-TEST-001", "SS-TEST-002", "SS-ACC-001", "SS-CERT-001", "SS-DATE-001", "SS-INS-001"}
# how a slot constrains a product in the (synthetic) market catalogue
SLOT_MARKET = {
    "efficacy": ("attr", "luminous_efficacy", "min"), "power_factor": ("attr", "power_factor", "min"),
    "cri": ("attr", "colour_rendering_index", "min"), "star_rating": ("attr", "star_rating", "min"),
    "air_delivery": ("attr", "air_delivery", "min"), "surge": ("attr", "surge_protection", "min"),
    "breaking_capacity": ("attr", "breaking_capacity", "min"), "module_efficiency": ("attr", "module_efficiency", "min"),
    "module_power": ("attr", "module_power", "min"), "sweep": ("attr", "sweep", "eq"),
    "cap_type": ("list", "caps", "in"), "ip": ("list", "ip_codes", "in"), "conductor_area": ("list", "sizes_mm2", "in"),
    "rated_current": ("list", "ratings_a", "in"), "curve": ("list", "curves", "in"), "sizes": ("list", "sizes_mm", "subset"),
}


def _fmt_date(d: date) -> str:
    return d.strftime("%d-%m-%Y")


def _fmt_inr(v: float) -> str:
    v = int(round(v))
    s = str(v)
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts + [tail])


def _val(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, list):
        return ", ".join(_val(x) for x in v)
    return str(v)


def _item_label_short(it: dict) -> str:
    return it["label"].split("(")[0].strip()


# ------------------------------------------------------------------------------------------ market guard
def competition(item_id: str, slots: Dict[str, object]) -> dict:
    prods = [p for p in catalog() if p.get("item") == item_id and p.get("bis_licence", True)]
    if not prods:
        return {"available": False}

    def ok(p, active):
        for sid in active:
            if sid not in SLOT_MARKET or slots.get(sid) in (None, ""):
                continue
            kind, key, op = SLOT_MARKET[sid]
            v = slots[sid]
            if kind == "attr":
                pv = (p.get("attributes") or {}).get(key)
                if pv is None:
                    continue
                if item_id == "led_lamp" and sid == "rated_power":
                    continue
                if op == "min" and pv < float(v) - 1e-9:
                    return False
                if op == "eq" and abs(pv - float(v)) > 1e-6:
                    return False
            else:
                lst = p.get(key)
                if lst is None:
                    continue
                if op == "in" and v not in lst and _val(v) not in [_val(x) for x in lst]:
                    return False
                if op == "subset" and not set(v).issubset(set(lst)):
                    return False
        # buyer-chosen wattage: a lamp family must offer that wattage
        if "rated_power" in slots and p.get("attributes", {}).get("rated_power") is not None and item_id in ("led_lamp", "led_streetlight"):
            if abs(p["attributes"]["rated_power"] - float(slots["rated_power"])) > float(slots["rated_power"]) * 0.1 + 1e-9:
                return False
        return True

    active = list(slots.keys())
    elig = [p for p in prods if ok(p, active)]
    mfrs = sorted({p["manufacturer"] for p in elig})
    base_mfrs = {p["manufacturer"] for p in prods}
    binding = []
    for sid in active:
        if sid not in SLOT_MARKET:
            continue
        without = {p["manufacturer"] for p in prods if ok(p, [a for a in active if a != sid])}
        if len(without) > len(mfrs):
            binding.append({"slot": sid, "manufacturers_without": len(without), "excluded": len(without) - len(mfrs)})
    binding.sort(key=lambda b: -b["excluded"])
    return {"available": True, "licensed_manufacturers": len(base_mfrs), "eligible_manufacturers": len(mfrs), "eligible_names": mfrs,
            "eligible_products": len(elig), "binding": binding, "ok": len(mfrs) >= MIN_COMPETITION,
            "note": "Synthetic demonstration catalogue — plug in GeM / BIS licence data for production."}


def competition_guard(item: dict, slots: Dict[str, object], sources: Dict[str, str]) -> Tuple[Dict[str, object], List[dict], dict]:
    """Relax *practice* defaults (never the buyer's stated choices) until at least MIN_COMPETITION manufacturers comply."""
    changes = []
    comp = competition(item["id"], slots)
    guard = 0
    while comp.get("available") and not comp["ok"] and guard < 6:
        guard += 1
        moved = False
        for b in comp["binding"]:
            sid = b["slot"]
            sl = next((s for s in item["slots"] if s["id"] == sid), None)
            if not sl or sl.get("basis") != "practice" or sources.get(sid) in ("stated", "answered"):
                continue
            opts = sorted(o for o in (sl.get("options") or []) if isinstance(o, (int, float)))
            cur = float(slots[sid])
            lower = [o for o in opts if o < cur]
            if not lower:
                continue
            new = lower[-1]
            changes.append({"slot": sid, "label": sl["label"], "from": cur, "to": new,
                            "reason": f"'{sl['label']} ≥ {_val(cur)}' left {comp['eligible_manufacturers']} of {comp['licensed_manufacturers']} licensed manufacturers; relaxed to {_val(new)} (GFR 2017 Rule 144: specifications must not unduly restrict competition)."})
            slots[sid] = new
            moved = True
            break
        if not moved:
            break
        comp = competition(item["id"], slots)
    if comp.get("available") and not comp["ok"]:
        stated = [b for b in comp["binding"] if sources.get(b["slot"]) in ("stated", "answered")]
        alts = []
        for b in stated:
            sl = next((s for s in item["slots"] if s["id"] == b["slot"]), None)
            if not sl:
                continue
            opts = [o for o in (sl.get("options") or []) if o != slots[b["slot"]]]
            for o in opts:
                trial = competition(item["id"], {**slots, b["slot"]: o})
                if trial.get("eligible_manufacturers", 0) > comp["eligible_manufacturers"]:
                    alts.append({"slot": b["slot"], "label": sl["label"], "to": o, "manufacturers": trial["eligible_manufacturers"]})
        comp["alternatives"] = sorted(alts, key=lambda a: -a["manufacturers"])[:4]
        comp["warning"] = (f"Only {comp['eligible_manufacturers']} manufacturer(s) can meet this specification"
                           + (f"; your choice of {', '.join(b['slot'] for b in stated)} is the binding constraint." if stated else "."))
    return slots, changes, comp


# ------------------------------------------------------------------------------------------ clause plan
class Plan:
    def __init__(self):
        self.sections: List[dict] = []

    def section(self, title: str, page: int) -> dict:
        s = {"no": len(self.sections) + 1, "title": title, "page": page, "clauses": []}
        self.sections.append(s)
        return s

    @staticmethod
    def clause(sec: dict, heading: str, text: str, basis: str, provenance: List[dict], item_id: Optional[str] = None, **kw) -> dict:
        c = {"key": f"{sec['no']}.{len(sec['clauses']) + 1}", "heading": heading, "text": text, "basis": basis,
             "provenance": provenance, "item_id": item_id, **kw}
        sec["clauses"].append(c)
        return c

    def all(self) -> List[dict]:
        return [c for s in self.sections for c in s["clauses"]]


def _render_template(kg: ProcurementKG, tpl: str, slots: Dict[str, object], on: date, item_id: str) -> Tuple[str, List[dict]]:
    prov = []

    def std(m):
        des = m.group(1).strip()
        r = kg.resolve_designation(des, on)
        prov.append({"kind": "standard", "designation": des, "cite_as": r["cite_as"], "status": r["status"],
                     "concurrent": r["concurrent_editions"], "amendments_in_force": len(r["amendments_in_force"]),
                     "path": [f"item:{item_id}", "GOVERNED_BY", f"fam:{family_of(des)}", "HAS_EDITION", f"std:{r['cite_as']}"],
                     "source_url": r.get("source_url") or "https://standards.bis.gov.in/"})
        return r["cite_as"]
    text = re.sub(r"\{std:([^}]+)\}", std, tpl)
    text = re.sub(r"\{(\w+)\}", lambda m: _val(slots.get(m.group(1), m.group(0))), text)
    return text, prov


def _estimate(lines: List[dict]) -> float:
    tot = 0.0
    for li in lines:
        it = item_by_id(li["item_id"])
        lo, hi = it.get("estimate_inr") or [0, 0]
        tot += (li["quantity"] or 0) * (lo + hi) / 2
    return tot


def build_plan(intent: dict, kg: ProcurementKG, on: date, meta: dict, overrides: Optional[Dict[str, Dict[str, object]]] = None) -> dict:
    lines = [li for li in intent["line_items"] if li.get("item_id") and li.get("quantity")]
    if not lines:
        raise ValueError("No resolvable line items — answer the open questions first.")
    glob = intent.get("globals") or {}
    pol = {p["key"]: p for p in policies()}
    bid_days = max(pol["GFR-161"].get("min_bid_days", 21), int(meta.get("bid_days") or 21))
    publication = on
    prebid = on + timedelta(days=7)
    bid_due = on + timedelta(days=bid_days)
    delivery_days = int(glob.get("delivery_days") or max(item_by_id(li["item_id"]).get("delivery_days", 30) for li in lines))
    place = glob.get("delivery_place") or meta.get("delivery_place") or "the consignee's store as specified in the Purchase Order"
    est = float(glob.get("budget_inr") or 0) or _estimate(lines)
    est_source = "budget stated by the officer" if glob.get("budget_inr") else "indicative estimate from synthetic price bands — replace with market survey / last purchase price"
    emd = max(100, round(est * pol["GFR-170"]["rate"] / 100) * 100)
    names = [_item_label_short(item_by_id(li["item_id"])) for li in lines]
    names = [n + ("s" if item_by_id(li["item_id"])["unit_kind"] == "count" and not n.endswith("s") else "") for n, li in zip(names, lines)]
    items_txt = names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1]
    purpose = glob.get("purpose") or meta.get("purpose")
    title = meta.get("title") or f"Supply of {items_txt}" + (f" for {purpose}" if purpose else "")
    org = meta.get("organisation") or "Demo Government Department"
    ref = meta.get("reference_no") or f"SS/{re.sub(r'[^A-Z]', '', org.upper())[:6] or 'DEPT'}/{on.year}/{hashlib.sha1((title + on.isoformat()).encode()).hexdigest()[:4].upper()}"

    plan = Plan()
    guard_log: List[dict] = []
    comp_by_line: Dict[int, dict] = {}
    watch: List[dict] = []
    reg_register: Dict[str, dict] = {}
    std_register: Dict[str, dict] = {}

    # ---- 1. general
    s1 = plan.section("General Information", 1)
    P = Plan.clause
    P(s1, "Name of Work", f"{title}.", "buyer", [{"kind": "intent", "detail": intent.get("input", "")[:200]}])
    P(s1, "Estimated Value", f"Rs. {_fmt_inr(est)} (inclusive of all taxes and duties).", "buyer",
      [{"kind": "estimate", "detail": est_source}])
    P(s1, "Date of Publication", _fmt_date(publication), "policy", [{"kind": "policy", "key": "GFR-161"}])
    P(s1, "Pre-bid Meeting", f"{_fmt_date(prebid)} at 11:00 hrs in the office of the Purchaser; written queries may be sent up to one day before the meeting.", "practice", [])
    P(s1, "Last Date of Bid Submission", f"{_fmt_date(bid_due)} at 15:00 hrs ({bid_days} days from publication).", "policy",
      [{"kind": "policy", "key": "GFR-161", "detail": pol["GFR-161"]["summary"]}])
    P(s1, "Bid Validity", "90 days from the last date of bid submission.", "practice", [])
    P(s1, "Bid Security (EMD)", pol["GFR-170"]["clause"]["text"].format(emd=_fmt_inr(emd)), "policy",
      [{"kind": "policy", "key": "GFR-170", "detail": pol["GFR-170"]["summary"], "source_url": pol["GFR-170"]["source_url"],
        "verification_state": pol["GFR-170"]["verification_state"]}])
    P(s1, "Performance Security", pol["GFR-171"]["clause"]["text"], "policy",
      [{"kind": "policy", "key": "GFR-171", "detail": pol["GFR-171"]["summary"], "source_url": pol["GFR-171"]["source_url"],
        "verification_state": pol["GFR-171"]["verification_state"]}])

    # ---- 2. schedule of requirements
    s2 = plan.section("Schedule of Requirements", 1)
    for li in lines:
        it = item_by_id(li["item_id"])
        qty = li["quantity"]
        unit = li.get("unit") or it["default_unit"]
        P(s2, f"Item {li['index']}", f"{it['label']} — quantity {_val(round(qty, 3))} {unit}, as specified in Section {2 + lines.index(li) + 1}.",
          "buyer", [{"kind": "intent", "detail": li.get("text")}], item_id=it["id"], hsn=it.get("hsn"))

    # ---- 3..n technical specifications per item
    for li in lines:
        it = item_by_id(li["item_id"])
        slots = dict(li.get("slots") or {})
        slots.update((overrides or {}).get(str(li["index"]), {}))
        sources = dict(li.get("slot_sources") or {})
        for k in (overrides or {}).get(str(li["index"]), {}):
            sources[k] = "answered"
        slots, changes, comp = competition_guard(it, slots, sources)
        for ch in changes:
            ch["line"] = li["index"]
        guard_log += changes
        comp_by_line[li["index"]] = comp
        li["resolved_slots"] = slots
        b = kg.bundle(it["id"], on, horizon_days=delivery_days + (bid_days) + 60)
        for s in b["standards"]:
            std_register[s["cite_as"]] = {"cite_as": s["cite_as"], "title": s.get("title"), "role": s["role"], "item": it["label"],
                                         "concurrent": s["concurrent_editions"], "amendments_in_force": [a["number"] for a in s["amendments_in_force"]],
                                         "timeline": s["timeline"]}
        for g in b["regulations_in_force"] + b["regulations_upcoming"]:
            reg_register[g["key"]] = {**{k: g.get(k) for k in ("key", "title", "scheme", "regulation_type", "issuing_authority", "effective_from",
                                                                  "verification_state", "source_url", "via")},
                                      "item": it["label"], "in_force": g in b["regulations_in_force"]}
        for w in b["watchlist"]:
            watch.append({**w, "item": it["label"]})
        sec = plan.section(f"Technical Specifications — Item {li['index']}: {it['label']}", 2)
        for tpl in it["spec"]:
            cond = tpl.get("when")
            if cond and slots.get(cond["slot"]) not in cond["in"]:
                continue
            text, prov = _render_template(kg, tpl["text"], slots, on, it["id"])
            basis = tpl["basis"]
            if tpl.get("regulation"):
                g = reg_register.get(tpl["regulation"]) or next((x for x in kg.kb.regulations if x.key == tpl["regulation"]), None)
                gd = g if isinstance(g, dict) else (g.to_dict() if g else {})
                prov.append({"kind": "regulation", "key": tpl["regulation"], "title": gd.get("title"), "scheme": gd.get("scheme"),
                             "verification_state": gd.get("verification_state"), "source_url": gd.get("source_url"),
                             "path": [f"item:{it['id']}", "GOVERNED_BY", f"fam:{gd.get('via') or '?'}", "◀MANDATES", f"reg:{tpl['regulation']}"]})
            used = re.findall(r"\{(\w+)\}", tpl["text"])
            for sid in used:
                sl = next((s for s in it["slots"] if s["id"] == sid), None)
                if sl:
                    prov.append({"kind": "slot", "slot": sid, "label": sl["label"], "value": _val(slots.get(sid)),
                                 "source": sources.get(sid, "default"), "basis": sl.get("basis"),
                                 "guarded": any(c["slot"] == sid and c["line"] == li["index"] for c in guard_log)})
            if basis == "buyer" and not any(sources.get(s) in ("stated", "answered") for s in used):
                basis = "buyer_default"
            P(sec, tpl["heading"], text, basis, prov, item_id=it["id"])
        # dependency closure: standards the product standard itself relies on (KG NORMATIVE_REF / SAFETY_REQ / TEST_METHOD)
        cited_now = " ".join(c["text"] for c in sec["clauses"])
        refs, rprov = [], []
        for s_ in b["standards"]:
            node = f"std:{s_['cite_as']}"
            if node not in kg.g:
                continue
            for _, v, k in kg.g.out_edges(node, keys=True):
                if k in ("NORMATIVE_REF", "SAFETY_REQ", "TEST_METHOD"):
                    lab = kg.g.nodes[v]["label"]
                    if lab.split(":")[0] in cited_now or lab in refs:
                        continue
                    refs.append(lab)
                    rprov.append({"kind": "standard", "designation": lab, "cite_as": lab, "status": kg.g.nodes[v].get("status", "current"),
                                  "path": [f"item:{it['id']}", "GOVERNED_BY", node, k, v], "relation": k})
        if refs:
            P(sec, "Referenced Standards", f"The following standards referred to in the governing specification shall also apply to the extent specified therein: {', '.join(refs)}.",
              "standard", rprov, item_id=it["id"])

    # ---- testing, acceptance, inspection
    sq = plan.section("Testing, Inspection and Acceptance", 3)
    tt_parts, tt_prov = [], []
    for li in lines:
        it = item_by_id(li["item_id"])
        type_tests = [t for t in it.get("tests", []) if t["stage"] == "type"]
        if type_tests:
            txt, prov = _render_template(kg, ", ".join(f"{t['name'].lower()} as per {{std:{t['standard']}}}" for t in type_tests), {}, on, it["id"])
            tt_parts.append(f"for {_item_label_short(it)}, {txt}")
            tt_prov += prov
    if tt_parts:
        P(sq, "Type Tests", "Type test reports for each offered model from a NABL-accredited laboratory having the tests in its scope of accreditation shall be submitted with the bid, covering "
          + "; ".join(tt_parts) + ". The reports shall not be older than three years on the last date of bid submission.", "standard", tt_prov)
    for li in lines:
        it = item_by_id(li["item_id"])
        acc = [t for t in it.get("tests", []) if t["stage"] == "acceptance"]
        if acc:
            txt, prov = _render_template(kg, "; ".join(f"{t['name']} (as per {{std:{t['standard']}}})" for t in acc), {}, on, it["id"])
            if it["unit_kind"] == "mass":
                sampling = "Sampling: samples shall be drawn from each lot as defined in the cited standard"
            elif it["unit_kind"] == "length":
                sampling = "Sampling: one sample coil shall be drawn from every twenty coils or part thereof"
            else:
                sampling = "Sampling: two samples shall be drawn at random by the Purchaser from every lot of up to one thousand units or part thereof"
            P(sq, f"Acceptance — {_item_label_short(it)}", f"{txt[0].upper() + txt[1:]} shall be carried out in a NABL-accredited laboratory on samples from each lot. {sampling}. "
              "A lot failing any acceptance test on re-test of twice the number of samples shall be rejected and replaced at the supplier's cost.",
              "practice", prov, item_id=it["id"])
    P(sq, "Inspection", "The Purchaser or its authorised representative shall carry out pre-dispatch inspection at the manufacturer's works on seven days' notice and receipt inspection at the consignee's store; the cost of testing shall be borne by the supplier.",
      "practice", [])

    # ---- delivery & warranty
    sd = plan.section("Delivery and Warranty", 3)
    P(sd, "Delivery Period", f"The entire quantity shall be delivered within {delivery_days} days from the date of issue of the Purchase Order.", "buyer",
      [{"kind": "intent", "detail": "stated" if glob.get("delivery_days") else "item default"}])
    P(sd, "Place of Delivery", f"Delivery shall be made at {place}, freight and insurance paid.", "buyer", [])
    for li in lines:
        it = item_by_id(li["item_id"])
        if it.get("warranty_text"):
            txt, prov = _render_template(kg, it["warranty_text"], li["resolved_slots"], on, it["id"])
            P(sd, f"Warranty — {_item_label_short(it)}", txt, "practice", prov, item_id=it["id"])

    # ---- eligibility & evaluation (policy layer)
    se = plan.section("Eligibility, Preferences and Evaluation", 3)
    for key in ("PPP-MII-2017", "MSE-PPP-2012"):
        p = pol[key]
        P(se, p["clause"]["heading"], p["clause"]["text"], "policy",
          [{"kind": "policy", "key": key, "detail": p["summary"], "source_url": p.get("source_url"), "verification_state": p.get("verification_state")}])
    P(se, pol["GFR-144"]["clause"]["heading"], pol["GFR-144"]["clause"]["text"], "policy",
      [{"kind": "policy", "key": "GFR-144", "detail": pol["GFR-144"]["summary"], "source_url": pol["GFR-144"]["source_url"]}])
    P(se, "Evaluation", "Bids shall first be evaluated for technical responsiveness against every clause of Sections 3 to "
      f"{2 + len(lines)}; the contract shall be awarded item-wise to the lowest-priced technically responsive bidder (L1).", "policy",
      [{"kind": "policy", "key": "GFR-144"}])

    return {"title": title, "reference_no": ref, "organisation": org, "publication": publication.isoformat(), "bid_due": bid_due.isoformat(),
            "prebid": prebid.isoformat(), "delivery_days": delivery_days, "estimated_value": est, "estimate_source": est_source, "emd": emd,
            "plan": plan, "lines": lines, "competition": comp_by_line, "guard": guard_log, "watchlist": sorted(watch, key=lambda w: w.get("date") or ""),
            "standards_register": list(std_register.values()), "regulations_register": list(reg_register.values())}


def render_text(meta: dict, plan: Plan) -> str:
    out = [f"--- Page 1 ---", meta["organisation"].upper(), f"NOTICE INVITING TENDER No. {meta['reference_no']}", meta["title"], ""]
    page = 1
    for s in plan.sections:
        if s["page"] != page:
            page = s["page"]
            out += ["", f"--- Page {page} ---"]
        out += ["", f"{s['no']}. {s['title']}"]
        for c in s["clauses"]:
            out.append(f"{c['key']} {c['heading']}: {c['text']}")
    return "\n".join(out) + "\n"


def _apply_fix(plan: Plan, fix: dict) -> bool:
    op = fix.get("op")
    target = next((c for c in plan.all() if c["key"] == fix.get("clause_key") or c["key"] == fix.get("after_key")), None)
    if op == "append" and target:
        if fix["text"] in target["text"]:
            return False
        target["text"] = target["text"].rstrip(".") + ". " + fix["text"]
        target.setdefault("repairs", []).append(fix["text"])
        return True
    if op == "add_clause":
        sec = next((s for s in plan.sections if target in s["clauses"]), plan.sections[-1])
        if any(fix["text"] in c["text"] for c in sec["clauses"]):
            return False
        heading, _, body = fix["text"].partition(":")
        Plan.clause(sec, heading.strip() if body else "Additional Requirement", (body or fix["text"]).strip(), "repair", [{"kind": "repair", "detail": "added by the self-verification loop"}])
        return True
    return False


def synthesize(intent: dict, kb: KnowledgeBase, on: Optional[date] = None, meta: Optional[dict] = None,
               overrides: Optional[Dict[str, Dict[str, object]]] = None, max_passes: int = 4) -> dict:
    on = on or date.today()
    meta = dict(meta or {})
    kg = get_kg(kb)
    res = build_plan(intent, kg, on, meta, overrides)
    plan: Plan = res["plan"]
    passes = []
    twin = None
    for i in range(max_passes):
        text = render_text(res, plan)
        ir = build_ir(parse_text(text, "generated_tender.txt"), doc_type="tender")
        twin = compile_ir(ir, kb, as_of=on, tender_meta={"title": res["title"], "reference_no": res["reference_no"], "organisation": res["organisation"]})
        crit = [f for f in twin["findings"] if f["severity"] in ("critical", "major")]
        fixed = []
        for f in crit:
            if f["rule_id"] in SAFE_FIX_RULES and f.get("fix") and _apply_fix(plan, f["fix"]):
                fixed.append({"rule_id": f["rule_id"], "clause_key": f.get("clause_key"), "message": f["message"]})
        passes.append({"pass": i + 1, "readiness": twin["readiness"]["score"], "gate": twin["readiness"]["gate"],
                       "counts": twin["readiness"]["counts"], "auto_repaired": fixed,
                       "open": [{"rule_id": f["rule_id"], "severity": f["severity"], "clause_key": f.get("clause_key"), "message": f["message"]} for f in crit
                                if not any(x["rule_id"] == f["rule_id"] and x["clause_key"] == f.get("clause_key") for x in fixed)]})
        if not fixed:
            break
    text = render_text(res, plan)
    clauses = plan.all()
    basis_counts: Dict[str, int] = {}
    for c in clauses:
        basis_counts[c["basis"]] = basis_counts.get(c["basis"], 0) + 1
    return {
        "generator": GENERATOR_VERSION, "as_of": on.isoformat(),
        "meta": {k: res[k] for k in ("title", "reference_no", "organisation", "publication", "bid_due", "prebid", "delivery_days",
                                     "estimated_value", "estimate_source", "emd")},
        "text": text, "sections": [{"no": s["no"], "title": s["title"], "clauses": s["clauses"]} for s in plan.sections],
        "basis_counts": basis_counts,
        "lines": [{"index": li["index"], "item_id": li["item_id"], "item_label": item_by_id(li["item_id"])["label"], "quantity": li["quantity"],
                   "unit": li.get("unit"), "slots": li["resolved_slots"], "slot_sources": li.get("slot_sources", {}),
                   "competition": res["competition"].get(li["index"])} for li in res["lines"]],
        "competition_guard": res["guard"], "watchlist": res["watchlist"],
        "standards_register": res["standards_register"], "regulations_register": res["regulations_register"],
        "verification": {"passes": passes, "final": {"readiness": twin["readiness"]["score"], "gate": twin["readiness"]["gate"],
                                                    "counts": twin["readiness"]["counts"]},
                         "open_findings": [{"rule_id": f["rule_id"], "severity": f["severity"], "clause_key": f.get("clause_key"), "message": f["message"],
                                            "suggestion": f.get("suggestion")} for f in twin["findings"]],
                         "proofs": {"conflicts": len(twin["verification"].get("conflicts", [])), "engine": twin["verification"].get("engine")}},
        "route_advice": route_advice(res["estimated_value"]),
        "kg_snapshot": kb.snapshot_hash()[:16],
    }


def route_advice(value: float) -> dict:
    """GFR 2017 procurement route for common-use goods (Rule 149 GeM; Rules 154/155/161). Thresholds per the curated policy layer."""
    if value <= 50000:
        route, rule = "Direct purchase on GeM from any seller meeting the requirement", "GFR 149(i)"
    elif value <= 1000000:
        route, rule = "GeM: L1 among at least three OEMs/sellers meeting the requirement (or GeM bid)", "GFR 149(ii)"
    else:
        route, rule = "GeM bid / reverse auction (mandatory above Rs 10 lakh for goods available on GeM)", "GFR 149(iii)"
    extra = "If the goods are not available on GeM, an advertised tender enquiry is required at Rs 50 lakh and above (Rule 161)." if value >= 5000000 else \
        "If not available on GeM, limited tender enquiry may be used below Rs 50 lakh (Rule 162)."
    return {"estimated_value": value, "route": route, "rule": rule, "note": extra,
            "caveat": "Thresholds from the curated policy layer (secondary source) — verify current DoE limits before use."}


# ------------------------------------------------------------------------------------------ DOCX export
def to_docx(gen: dict) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    m = gen["meta"]
    p = d.add_paragraph(m["organisation"].upper())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p = d.add_paragraph(f"NOTICE INVITING TENDER No. {m['reference_no']}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h = d.add_heading(m["title"], level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = d.add_paragraph(f"Draft generated by SpecSure Tender Studio ({gen['generator']}) as of {gen['as_of']}; knowledge-graph snapshot {gen['kg_snapshot']}. "
                           f"Compiler verification: readiness {gen['verification']['final']['readiness']}, gate {gen['verification']['final']['gate']}. "
                           "This is decision support — the competent authority must review every clause before publication.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(8.5)
    note.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    for s in gen["sections"]:
        d.add_heading(f"{s['no']}. {s['title']}", level=2)
        for c in s["clauses"]:
            para = d.add_paragraph()
            r = para.add_run(f"{c['key']} {c['heading']}: ")
            r.bold = True
            para.add_run(c["text"])
    d.add_page_break()
    d.add_heading("Annex T — Clause traceability (why each clause exists)", level=2)
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for i, hname in enumerate(["Clause", "Basis", "Source (knowledge graph)", "Verification"]):
        t.rows[0].cells[i].text = hname
    for s in gen["sections"]:
        for c in s["clauses"]:
            src, ver = [], []
            for pv in c.get("provenance", []):
                if pv["kind"] == "standard":
                    src.append(f"{pv['cite_as']} (edition in force)")
                elif pv["kind"] == "regulation":
                    src.append(pv.get("title") or pv["key"])
                    ver.append(pv.get("verification_state") or "")
                elif pv["kind"] == "policy":
                    src.append(pv["key"])
                    ver.append(pv.get("verification_state") or "")
                elif pv["kind"] == "slot":
                    src.append(f"{pv['label']} = {pv['value']} ({pv['source']})")
            row = t.add_row().cells
            row[0].text, row[1].text, row[2].text, row[3].text = c["key"], c["basis"], "; ".join(src)[:400], ", ".join(x for x in ver if x)
    d.add_heading("Annex S — Standards & regulations register (as of the publication date)", level=2)
    for s in gen["standards_register"]:
        extra = f" — concurrent edition(s): {', '.join(x['is_number'] + ' until ' + (x['valid_to'] or '?') for x in s['concurrent'])}" if s["concurrent"] else ""
        d.add_paragraph(f"{s['cite_as']} — {s.get('title') or ''} ({s['role']}; {s['item']}){extra}", style="List Bullet")
    for g in gen["regulations_register"]:
        d.add_paragraph(f"{g['title']} — {g.get('scheme') or ''}; {'in force' if g['in_force'] else 'upcoming ' + str(g.get('effective_from'))}; verification: {g.get('verification_state')}",
                        style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
