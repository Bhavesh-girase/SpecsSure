"""
Layout-aware, page-preserving document ingestion with a tiered strategy:

  1. Docling (IBM, open-source) when installed and SPECSURE_USE_DOCLING=1
     -> best layout + table structure recovery
  2. PyMuPDF text layer + PyMuPDF table finder (fast, default)
  3. OCR for scanned pages / images: RapidOCR PP-OCRv6 (+ PP-OCRv5 Devanagari) with domain-aware repair
     (engine/ocr.py); Tesseract only as last resort

Every page keeps its page number so that each extracted requirement, finding and
piece of evidence can point back to "document X, page N, clause Y".
"""
from __future__ import annotations

import hashlib
import io
import logging
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional

logger = logging.getLogger(__name__)

PAGE_MARKER = re.compile(r"^-{2,}\s*Page\s+(\d+)\s*-{2,}\s*$", re.IGNORECASE | re.MULTILINE)


@dataclass
class PageText:
    page_no: int
    text: str
    tables: List[List[List[str]]] = field(default_factory=list)
    ocr: bool = False
    ocr_info: Optional[dict] = None


@dataclass
class ParsedDocument:
    filename: str
    sha256: str
    pages: List[PageText]
    parser: str
    ocr_pages: List[int] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(f"--- Page {p.page_no} ---\n{p.text}" for p in self.pages)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def meta(self) -> dict:
        ocr = [{"page": p.page_no, **{k: p.ocr_info[k] for k in ("engine", "mean_conf", "ms") if k in p.ocr_info},
                "repairs": p.ocr_info.get("repairs", [])[:10], "languages": p.ocr_info.get("languages")} for p in self.pages if p.ocr_info]
        return {"filename": self.filename, "sha256": self.sha256, "parser": self.parser,
                "page_count": self.page_count, "ocr_pages": self.ocr_pages, "warnings": self.warnings,
                "tables": sum(len(p.tables) for p in self.pages), "ocr": ocr}


def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def parse_document(data: bytes, filename: str) -> ParsedDocument:
    name = (filename or "document.txt").lower()
    if data[:5] == b"%PDF-" and not name.endswith(".pdf"):
        name += ".pdf"                         # sniff: bid-document URLs often have no extension
    elif data[:2] == b"PK" and not name.endswith(".docx") and b"word/" in data[:4000]:
        name += ".docx"
    digest = sha256_bytes(data)
    try:
        if name.endswith(".pdf"):
            return _parse_pdf(data, filename, digest)
        if name.endswith(".docx"):
            return _parse_docx(data, filename, digest)
        if name.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")):
            from app.engine.ocr import ocr_image
            r = ocr_image(data)
            warn = [] if r.text.strip() else ["OCR engine unavailable or produced no text"]
            if r.text.strip() and r.mean_conf and r.mean_conf < 0.8:
                warn.append(f"Low OCR confidence ({r.mean_conf:.0%}) — verify extracted values against the image")
            info = {"engine": r.engine, "mean_conf": r.mean_conf, "ms": r.ms, "repairs": r.repairs, "languages": r.languages}
            pages = [PageText(1, r.text, _tables_from_pipe_text(r.text), ocr=True, ocr_info=info)]
            return ParsedDocument(filename, digest, pages, "ocr-image", ocr_pages=[1], warnings=warn)
        return parse_text(data.decode("utf-8", errors="replace"), filename, digest)
    except Exception as e:  # never crash ingestion; surface a warning instead
        logger.exception("Document parse failed for %s", filename)
        return ParsedDocument(filename, digest, [PageText(1, "")], "failed", warnings=[f"Parse failed: {e}"])


def parse_text(text: str, filename: str = "input.txt", digest: Optional[str] = None) -> ParsedDocument:
    digest = digest or sha256_bytes(text.encode("utf-8"))
    text = text.replace("\r\n", "\n")
    pages: List[PageText] = []
    if PAGE_MARKER.search(text):
        parts = PAGE_MARKER.split(text)
        # parts = [pre, n1, t1, n2, t2 ...]
        if parts[0].strip():
            pages.append(PageText(1, parts[0].strip("\n")))
        for i in range(1, len(parts) - 1, 2):
            pages.append(PageText(int(parts[i]), parts[i + 1].strip("\n")))
    elif "\f" in text:
        for i, chunk in enumerate(text.split("\f")):
            pages.append(PageText(i + 1, chunk.strip("\n")))
    else:
        pages.append(PageText(1, text.strip("\n")))
    for p in pages:
        p.tables = _tables_from_pipe_text(p.text)
    return ParsedDocument(filename, digest, pages, "text")


def _tables_from_pipe_text(text: str) -> List[List[List[str]]]:
    tables, current = [], []
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("|") and s.count("|") >= 2:
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            current.append(cells)
        else:
            if len(current) >= 2:
                tables.append(current)
            current = []
    if len(current) >= 2:
        tables.append(current)
    return tables


def _parse_pdf(data: bytes, filename: str, digest: str) -> ParsedDocument:
    if os.getenv("SPECSURE_USE_DOCLING") == "1":
        try:
            return _parse_with_docling(data, filename, digest)
        except Exception as e:
            logger.warning("Docling unavailable/failed (%s); falling back to PyMuPDF", e)
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    pages: List[PageText] = []
    ocr_pages: List[int] = []
    warnings: List[str] = []
    for i, page in enumerate(doc):
        text = page.get_text("text") or ""
        tables: List[List[List[str]]] = []
        try:
            tf = page.find_tables()
            for t in tf.tables:
                rows = [[(c or "").strip() for c in row] for row in t.extract()]
                if len(rows) >= 2:
                    tables.append(rows)
        except Exception:
            pass
        is_ocr = False
        info = None
        if len(text.strip()) < 25 and (page.get_images() or not text.strip()):
            ocr_text, info = _ocr_pdf_page(page)
            if ocr_text:
                text, is_ocr = ocr_text, True
                ocr_pages.append(i + 1)
                if info and info.get("mean_conf") and info["mean_conf"] < 0.8:
                    warnings.append(f"Page {i + 1}: low OCR confidence ({info['mean_conf']:.0%})")
            else:
                warnings.append(f"Page {i + 1} appears scanned but no OCR engine is available")
        # render tables as pipe rows so downstream clause parsing sees them in-line
        if tables and "|" not in text:
            for t in tables:
                text += "\n" + "\n".join("| " + " | ".join(r) + " |" for r in t)
        pages.append(PageText(i + 1, text, tables, is_ocr, info))
    doc.close()
    return ParsedDocument(filename, digest, pages, "pymupdf", ocr_pages, warnings)


def _parse_with_docling(data: bytes, filename: str, digest: str) -> ParsedDocument:
    import tempfile
    from docling.document_converter import DocumentConverter  # type: ignore

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(data)
        path = tmp.name
    result = DocumentConverter().convert(path)
    md = result.document.export_to_markdown(page_break_placeholder="\f")
    parsed = parse_text(md, filename, digest)
    parsed.parser = "docling"
    return parsed


def _ocr_pdf_page(page):
    """Render the page and run the OCR router (RapidOCR PP-OCRv6 → Devanagari → Tesseract)."""
    try:
        from app.engine.ocr import ocr_pdf_page
        r = ocr_pdf_page(page, dpi=200)
        if r.text.strip():
            return r.text, {"engine": r.engine, "mean_conf": r.mean_conf, "ms": r.ms, "repairs": r.repairs, "languages": r.languages}
    except Exception as e:
        logger.warning("OCR router failed: %s", e)
    try:  # PyMuPDF's built-in Tesseract bridge
        tp = page.get_textpage_ocr(language="eng", dpi=300, full=True)
        txt = page.get_text("text", textpage=tp)
        if txt.strip():
            return txt, {"engine": "tesseract (pymupdf)", "mean_conf": None, "repairs": []}
    except Exception:
        pass
    return "", None


def _ocr_image_bytes(img: bytes) -> str:
    from app.engine.ocr import ocr_image
    return ocr_image(img).text


def _parse_docx(data: bytes, filename: str, digest: str) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    lines: List[str] = []
    tables: List[List[List[str]]] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") and not re.match(r"^\d", t):
                t = t.upper()
            lines.append(t)
        elif tag == "tbl":
            tbl = Table(child, d)
            rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
            tables.append(rows)
            for r in rows:
                lines.append("| " + " | ".join(r) + " |")
    text = "\n".join(lines)
    parsed = parse_text(text, filename, digest)
    parsed.parser = "python-docx"
    return parsed
